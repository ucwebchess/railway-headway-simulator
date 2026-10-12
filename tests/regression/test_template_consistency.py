"""Regression tests for Excel template consistency, canonical enum alignment, and cross-workbook integrity.

Addresses RHS-AUDIT-TEMPLATES-001 and satisfies RHS-CORRECTION-TEMPLATES-001:
- Verifies that all 6 shipped templates load and validate with zero errors.
- Verifies that example rows include ND_02, SIG_02, and SCN_ALT_01 to close the reference graph.
- Verifies that SharedResourceGroups.resource_ids references existing resources (and is validated).
- Verifies that SignallingSystem.TECHNOLOGY_TYPE allowed values match all canonical enum members.
- Verifies that the INSTRUCTIONS worksheet reflects updated baseline demonstration wording.
"""

from pathlib import Path
import openpyxl
import pytest

from headway.data.canonical import SignallingTechnologyType
from headway.data.converter import convert_raw_to_canonical
from headway.data.excel_registry import ALL_WORKBOOK_DEFS
from headway.data.importer import ExcelImporter
from headway.data.validation import Severity, ValidationReport
from headway.data.validator import DatasetValidator


TEMPLATE_DIR = Path("templates")
TEMPLATE_FILES = [
    "HEADWAY_INFRASTRUCTURE_v1.0.xlsx",
    "HEADWAY_ROLLING_STOCK_v1.0.xlsx",
    "HEADWAY_SIGNALLING_v1.0.xlsx",
    "HEADWAY_OPERATIONS_v1.0.xlsx",
    "HEADWAY_ANALYSIS_v1.0.xlsx",
    "HEADWAY_SCENARIOS_v1.0.xlsx",
]


@pytest.fixture
def loaded_templates():
    """Load all 6 template workbooks from templates/."""
    importer = ExcelImporter()
    report = ValidationReport()
    raw_wbs = {}
    for filename in TEMPLATE_FILES:
        path = TEMPLATE_DIR / filename
        assert path.exists(), f"Template file '{filename}' does not exist in templates/."
        rwb = importer.load_workbook(path, report)
        assert rwb is not None, f"Failed to load workbook '{filename}': {report.findings}"
        raw_wbs[rwb.workbook_type] = rwb

    return raw_wbs, report


def test_templates_load_with_zero_validation_errors(loaded_templates):
    """Assert that loading all six templates from templates/ returns zero validation errors."""
    raw_wbs, report = loaded_templates
    assert len(raw_wbs) == 6

    validator = DatasetValidator(report)
    val_report = validator.validate_all(raw_wbs)

    # Filter for ERROR or CRITICAL findings
    error_findings = [
        f for f in val_report.findings
        if f.severity in {Severity.ERROR, Severity.CRITICAL}
    ]
    assert not error_findings, f"Expected 0 validation errors, got {len(error_findings)}: {error_findings}"
    assert not val_report.has_errors

    # Project conversion should succeed cleanly
    project = convert_raw_to_canonical(raw_wbs)
    assert project is not None
    assert project.project_id == "PRJ_DEMO"


def test_nodes_example_contains_nd_02(loaded_templates):
    """Assert that Nodes worksheet example data contains ND_02."""
    raw_wbs, _ = loaded_templates
    infra = raw_wbs["INFRASTRUCTURE"]
    node_rows = infra.get_rows("Nodes")
    node_ids = [r["node_id"] for r in node_rows]

    assert "ND_01" in node_ids
    assert "ND_02" in node_ids, "Nodes example must contain ND_02."

    nd_02 = next(r for r in node_rows if r["node_id"] == "ND_02")
    assert nd_02["description"] == "Downstream endpoint"
    assert nd_02["node_type"] == "ENDPOINT"

    # Also assert in excel_registry
    infra_def = ALL_WORKBOOK_DEFS["INFRASTRUCTURE"]
    nodes_def = next(w for w in infra_def.worksheets if w.name == "Nodes")
    assert any(row[0] == "ND_02" for row in nodes_def.example_rows)


def test_signals_example_contains_sig_02(loaded_templates):
    """Assert that Signals worksheet example data contains SIG_02."""
    raw_wbs, _ = loaded_templates
    sig = raw_wbs["SIGNALLING"]
    signal_rows = sig.get_rows("Signals")
    signal_ids = [r["signal_id"] for r in signal_rows]

    assert "SIG_01" in signal_ids
    assert "SIG_02" in signal_ids, "Signals example must contain SIG_02."

    sig_02 = next(r for r in signal_rows if r["signal_id"] == "SIG_02")
    assert sig_02["link_id"] == "LNK_01"
    assert sig_02["offset_m"] == 1100.0
    assert sig_02["direction"] == "NOMINAL"
    assert sig_02["signal_type"] == "MAIN"

    # Also assert in excel_registry
    sig_def = ALL_WORKBOOK_DEFS["SIGNALLING"]
    signals_def = next(w for w in sig_def.worksheets if w.name == "Signals")
    assert any(row[0] == "SIG_02" for row in signals_def.example_rows)


def test_shared_resource_groups_resource_ids_reference_existing_resources(loaded_templates):
    """Assert that SharedResourceGroups.resource_ids reference existing resources."""
    raw_wbs, _ = loaded_templates
    infra = raw_wbs["INFRASTRUCTURE"]
    sig = raw_wbs["SIGNALLING"]

    srg_rows = infra.get_rows("SharedResourceGroups")
    assert len(srg_rows) >= 1
    srg = srg_rows[0]

    # Verify resource_ids in example data
    res_ids = srg["resource_ids"] if isinstance(srg["resource_ids"], list) else [srg["resource_ids"]]
    assert "BLK_01" in res_ids
    assert "BLK_02" not in res_ids, "BLK_02 should no longer be referenced unless defined."

    # Verify that all referenced IDs actually exist in Blocks or TVS
    block_ids = {r["block_id"] for r in sig.get_rows("Blocks")}
    tvs_ids = {r["tvs_id"] for r in infra.get_rows("TVSSections")}
    valid_ids = block_ids | tvs_ids

    for rid in res_ids:
        assert rid in valid_ids, f"Resource ID '{rid}' in SharedResourceGroups does not exist in Blocks or TVS."


def test_shared_resource_groups_validator_catches_missing_resource(loaded_templates):
    """Verify that DatasetValidator flags non-existent resource IDs in SharedResourceGroups."""
    raw_wbs, _ = loaded_templates
    report = ValidationReport()

    # Inject an invalid resource ID
    infra = raw_wbs["INFRASTRUCTURE"]
    infra.worksheets["SharedResourceGroups"] = [
        {
            "group_id": "GRP_TEST",
            "description": "Test conflict group",
            "resource_ids": ["NONEXISTENT_RES_99"],
        }
    ]

    validator = DatasetValidator(report)
    validator.validate_all(raw_wbs)

    assert report.has_errors
    srg_findings = [
        f for f in report.findings
        if f.worksheet == "SharedResourceGroups" and f.error_code == "ERR_REF_MISSING"
    ]
    assert len(srg_findings) >= 1
    assert "NONEXISTENT_RES_99" in srg_findings[0].message


def test_scenario_definitions_example_contains_scn_alt_01(loaded_templates):
    """Assert that ScenarioDefinitions worksheet example data contains SCN_ALT_01."""
    raw_wbs, _ = loaded_templates
    scn = raw_wbs["SCENARIOS"]
    scenario_rows = scn.get_rows("ScenarioDefinitions")
    scenario_ids = [r["scenario_id"] for r in scenario_rows]

    assert "SCN_BASELINE" in scenario_ids
    assert "SCN_ALT_01" in scenario_ids, "ScenarioDefinitions example must contain SCN_ALT_01."

    scn_alt = next(r for r in scenario_rows if r["scenario_id"] == "SCN_ALT_01")
    assert scn_alt["is_baseline"] is False
    assert "Alternative headway sensitivity" in scn_alt["description"]

    # Also assert in excel_registry
    scn_def = ALL_WORKBOOK_DEFS["SCENARIOS"]
    scenarios_def = next(w for w in scn_def.worksheets if w.name == "ScenarioDefinitions")
    assert any(row[0] == "SCN_ALT_01" for row in scenarios_def.example_rows)


def test_signalling_technology_type_allowed_values_include_all_canonical_enums():
    """Assert that SignallingSystem.TECHNOLOGY_TYPE allowed values include all five canonical enum members."""
    sig_def = ALL_WORKBOOK_DEFS["SIGNALLING"]
    ws_def = next(w for w in sig_def.worksheets if w.name == "SignallingSystem")
    tech_col = next(c for c in ws_def.columns if c.name == "TECHNOLOGY_TYPE")

    canonical_enums = {e.value for e in SignallingTechnologyType}
    registry_allowed = set(tech_col.allowed_values or [])

    assert len(canonical_enums) == 5
    assert registry_allowed == canonical_enums, (
        f"Registry allowed values {registry_allowed} do not match canonical enum members {canonical_enums}."
    )


def test_instructions_sheet_updated_wording():
    """Verify that the INSTRUCTIONS sheet in all six workbooks reflects the updated wording."""
    expected_phrase = (
        "Row 3 contains a minimal valid baseline network for demonstration purposes. "
        "You may modify or replace it with your own railway data. "
        "The example is designed to load without validation errors."
    )

    for filename in TEMPLATE_FILES:
        path = TEMPLATE_DIR / filename
        wb = openpyxl.load_workbook(path, data_only=True)
        assert "INSTRUCTIONS" in wb.sheetnames, f"Missing INSTRUCTIONS in {filename}"
        ws = wb["INSTRUCTIONS"]

        found = False
        for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            val = row[0]
            if val and expected_phrase in str(val):
                found = True
                break

        assert found, f"Template {filename} INSTRUCTIONS sheet does not contain updated wording."
