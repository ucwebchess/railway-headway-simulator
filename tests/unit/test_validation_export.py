"""Unit test for exporting validation findings to an Excel workbook."""

from pathlib import Path
import openpyxl
import pytest
from headway.data.validation import Severity, ValidationReport


@pytest.mark.unit
def test_export_validation_report_to_excel(tmp_path: Path):
    report = ValidationReport(project_id="PRJ_TEST_VAL")
    report.add_finding(
        error_code="ERR_TEST_01",
        severity=Severity.ERROR,
        message="Test error message",
        workbook="01_Infrastructure.xlsx",
        worksheet="TrackLinks",
        row=5,
        column="LENGTH_M",
        object_id="LNK_01",
        recommendation="Fix length",
    )
    report.add_finding(
        error_code="WARN_TEST_01",
        severity=Severity.WARNING,
        message="Test warning",
        workbook="01_Infrastructure.xlsx",
        worksheet="Platforms",
    )

    out_file = tmp_path / "findings.xlsx"
    exported_path = report.export_to_excel(out_file)

    assert exported_path.exists()
    wb = openpyxl.load_workbook(exported_path)
    assert "Validation_Findings" in wb.sheetnames
    ws = wb["Validation_Findings"]
    assert ws.max_row == 3  # Header + 2 findings
    assert ws.cell(row=2, column=1).value == "ERR_TEST_01"
    assert ws.cell(row=2, column=2).value == "ERROR"
    assert ws.cell(row=3, column=1).value == "WARN_TEST_01"
