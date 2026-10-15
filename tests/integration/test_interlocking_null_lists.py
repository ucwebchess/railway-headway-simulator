"""Integration and unit tests for empty list cells in InterlockingRoutes and comma-separated list fields.

Strictly verifies RHS-CORRECTION-INTERLOCKING-NULL-001:
- Empty, None, and NaN cells in optional comma-separated list fields parse as empty lists [].
- InterlockingRoutes.PROTECTED_BLOCKS and CONFLICTING_ROUTES do not raise 'NoneType' object is not iterable.
- Comma-separated token splitting strips whitespace consistently.
- Existing populated routes and list fields continue to parse accurately.
- User's exact failing scenario (RTE_01 with 447 links and empty protected blocks / conflicting routes)
  parses, validates, and converts to canonical model without error.
"""

from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd
import pytest

from headway.data.converter import convert_raw_to_canonical
from headway.data.excel_registry import ColumnDef, FieldType
from headway.data.importer import ExcelImporter, RawWorkbookData, _parse_cell_value, _parse_list_cell
from headway.data.validation import ValidationReport
from headway.data.validator import DatasetValidator


@pytest.mark.unit
def test_parse_list_cell_helper():
    """Verify _parse_list_cell helper handles None, NaN, whitespace, and tokens."""
    # None and NaN representations
    assert _parse_list_cell(None) == []
    assert _parse_list_cell(float("nan")) == []
    assert _parse_list_cell(np.nan) == []
    assert _parse_list_cell(pd.NA) == []

    # Empty strings and whitespace
    assert _parse_list_cell("") == []
    assert _parse_list_cell("   ") == []
    assert _parse_list_cell("\t\n") == []

    # Whitespace handling around commas
    assert _parse_list_cell("BLK_01, BLK_02") == ["BLK_01", "BLK_02"]
    assert _parse_list_cell(" BLK_01 , BLK_02 ") == ["BLK_01", "BLK_02"]
    assert _parse_list_cell("  BLK_01  ,   BLK_02  , BLK_03  ") == ["BLK_01", "BLK_02", "BLK_03"]

    # Semicolon tolerance
    assert _parse_list_cell("BLK_01; BLK_02") == ["BLK_01", "BLK_02"]

    # Existing list passthrough
    assert _parse_list_cell(["BLK_01", "BLK_02"]) == ["BLK_01", "BLK_02"]
    assert _parse_list_cell([]) == []


@pytest.mark.unit
def test_parse_cell_value_string_list_empty_and_nan():
    """Verify _parse_cell_value converts empty, None, and NaN cells to [] for optional STRING_LIST."""
    report = ValidationReport()
    loc_ctx = {"workbook": "test.xlsx", "worksheet": "InterlockingRoutes", "row": 2}

    col_opt = ColumnDef(
        name="PROTECTED_BLOCKS",
        canonical_name="protected_blocks",
        field_type=FieldType.STRING_LIST,
        required=False,
    )
    col_req = ColumnDef(
        name="LINK_SEQUENCE",
        canonical_name="link_sequence",
        field_type=FieldType.STRING_LIST,
        required=True,
    )

    # Optional empty cell -> []
    val, ok = _parse_cell_value(None, col_opt, report, loc_ctx)
    assert ok is True
    assert val == []

    # Optional NaN cell -> []
    val, ok = _parse_cell_value(float("nan"), col_opt, report, loc_ctx)
    assert ok is True
    assert val == []

    # Optional empty string -> []
    val, ok = _parse_cell_value("", col_opt, report, loc_ctx)
    assert ok is True
    assert val == []

    # Optional whitespace -> []
    val, ok = _parse_cell_value("   ", col_opt, report, loc_ctx)
    assert ok is True
    assert val == []

    # Required empty cell -> None with ERR_FIELD_REQUIRED_MISSING
    rep_req = ValidationReport()
    val, ok = _parse_cell_value(None, col_req, rep_req, loc_ctx)
    assert ok is False
    assert val is None
    assert any(f.error_code == "ERR_FIELD_REQUIRED_MISSING" for f in rep_req.findings)

    # Required NaN cell -> None with ERR_FIELD_REQUIRED_MISSING
    rep_req_nan = ValidationReport()
    val, ok = _parse_cell_value(float("nan"), col_req, rep_req_nan, loc_ctx)
    assert ok is False
    assert val is None
    assert any(f.error_code == "ERR_FIELD_REQUIRED_MISSING" for f in rep_req_nan.findings)


@pytest.mark.integration
def test_empty_protected_blocks_workbook_ingestion(tmp_path: Path):
    """Verify route with populated LINK_SEQUENCE and empty PROTECTED_BLOCKS parses without error."""
    sig_file = tmp_path / "02_Signalling.xlsx"
    wb = openpyxl.Workbook()

    # METADATA
    meta = wb.active
    meta.title = "METADATA"
    meta.append(["PROPERTY", "VALUE"])
    meta.append(["WORKBOOK_TYPE", "SIGNALLING"])
    meta.append(["SCHEMA_VERSION", "1.0.0"])

    # SignallingSystem
    ws_sys = wb.create_sheet("SignallingSystem")
    ws_sys.append(["TECHNOLOGY_TYPE", "ASPECT_MODEL", "DEFAULT_OVERLAP_M", "SIGHTING_TIME_S", "ROUTE_SETUP_TIME_S", "RELEASE_DELAY_S"])
    ws_sys.append(["GENERIC_FIXED_BLOCK_ENGINEERING_MODEL", "THREE_ASPECT", 50.0, 8.0, 3.0, 4.0])

    # Signals
    ws_sig = wb.create_sheet("Signals")
    ws_sig.append(["SIGNAL_ID", "LINK_ID", "OFFSET_M", "DIRECTION", "SIGNAL_TYPE"])
    ws_sig.append(["SIG_ENTRY", "LNK_01", 100.0, "NOMINAL", "MAIN"])
    ws_sig.append(["SIG_EXIT", "LNK_02", 500.0, "NOMINAL", "MAIN"])

    # Blocks
    ws_blk = wb.create_sheet("Blocks")
    ws_blk.append(["BLOCK_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID"])
    ws_blk.append(["BLK_01", "LNK_01", 0.0, 1000.0, "SIG_ENTRY", "SIG_EXIT"])

    # InterlockingRoutes: PROTECTED_BLOCKS is empty (None in Excel cell)
    ws_rte = wb.create_sheet("InterlockingRoutes")
    ws_rte.append(["ROUTE_ID", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID", "LINK_SEQUENCE", "PROTECTED_BLOCKS", "CONFLICTING_ROUTES"])
    ws_rte.append(["RTE_01", "SIG_ENTRY", "SIG_EXIT", "LNK_01, LNK_02", None, "RTE_02"])

    wb.save(sig_file)

    importer = ExcelImporter()
    report = ValidationReport()
    raw_wb = importer.load_workbook(sig_file, report)

    assert raw_wb is not None
    assert not report.has_critical
    routes = raw_wb.get_rows("InterlockingRoutes")
    assert len(routes) == 1
    route = routes[0]
    assert route["route_id"] == "RTE_01"
    assert route["link_sequence"] == ["LNK_01", "LNK_02"]
    assert route["protected_blocks"] == []
    assert route["conflicting_routes"] == ["RTE_02"]


@pytest.mark.integration
def test_empty_conflicting_routes_workbook_ingestion(tmp_path: Path):
    """Verify route with populated LINK_SEQUENCE and empty CONFLICTING_ROUTES parses without error."""
    sig_file = tmp_path / "02_Signalling.xlsx"
    wb = openpyxl.Workbook()

    meta = wb.active
    meta.title = "METADATA"
    meta.append(["PROPERTY", "VALUE"])
    meta.append(["WORKBOOK_TYPE", "SIGNALLING"])
    meta.append(["SCHEMA_VERSION", "1.0.0"])

    ws_sys = wb.create_sheet("SignallingSystem")
    ws_sys.append(["TECHNOLOGY_TYPE", "ASPECT_MODEL", "DEFAULT_OVERLAP_M", "SIGHTING_TIME_S", "ROUTE_SETUP_TIME_S", "RELEASE_DELAY_S"])
    ws_sys.append(["GENERIC_FIXED_BLOCK_ENGINEERING_MODEL", "THREE_ASPECT", 50.0, 8.0, 3.0, 4.0])

    ws_sig = wb.create_sheet("Signals")
    ws_sig.append(["SIGNAL_ID", "LINK_ID", "OFFSET_M", "DIRECTION", "SIGNAL_TYPE"])
    ws_sig.append(["SIG_ENTRY", "LNK_01", 100.0, "NOMINAL", "MAIN"])
    ws_sig.append(["SIG_EXIT", "LNK_02", 500.0, "NOMINAL", "MAIN"])

    ws_blk = wb.create_sheet("Blocks")
    ws_blk.append(["BLOCK_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID"])
    ws_blk.append(["BLK_01", "LNK_01", 0.0, 1000.0, "SIG_ENTRY", "SIG_EXIT"])

    # InterlockingRoutes: CONFLICTING_ROUTES is empty
    ws_rte = wb.create_sheet("InterlockingRoutes")
    ws_rte.append(["ROUTE_ID", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID", "LINK_SEQUENCE", "PROTECTED_BLOCKS", "CONFLICTING_ROUTES"])
    ws_rte.append(["RTE_01", "SIG_ENTRY", "SIG_EXIT", "LNK_01, LNK_02", "BLK_01", None])

    wb.save(sig_file)

    importer = ExcelImporter()
    report = ValidationReport()
    raw_wb = importer.load_workbook(sig_file, report)

    assert raw_wb is not None
    assert not report.has_critical
    routes = raw_wb.get_rows("InterlockingRoutes")
    assert len(routes) == 1
    route = routes[0]
    assert route["protected_blocks"] == ["BLK_01"]
    assert route["conflicting_routes"] == []


@pytest.mark.integration
def test_both_fields_empty_and_whitespace_handling(tmp_path: Path):
    """Verify route with both PROTECTED_BLOCKS and CONFLICTING_ROUTES empty parses without error,
    and verify whitespace handling in token lists."""
    sig_file = tmp_path / "02_Signalling.xlsx"
    wb = openpyxl.Workbook()

    meta = wb.active
    meta.title = "METADATA"
    meta.append(["PROPERTY", "VALUE"])
    meta.append(["WORKBOOK_TYPE", "SIGNALLING"])
    meta.append(["SCHEMA_VERSION", "1.0.0"])

    ws_sys = wb.create_sheet("SignallingSystem")
    ws_sys.append(["TECHNOLOGY_TYPE", "ASPECT_MODEL", "DEFAULT_OVERLAP_M", "SIGHTING_TIME_S", "ROUTE_SETUP_TIME_S", "RELEASE_DELAY_S"])
    ws_sys.append(["GENERIC_FIXED_BLOCK_ENGINEERING_MODEL", "THREE_ASPECT", 50.0, 8.0, 3.0, 4.0])

    ws_sig = wb.create_sheet("Signals")
    ws_sig.append(["SIGNAL_ID", "LINK_ID", "OFFSET_M", "DIRECTION", "SIGNAL_TYPE"])
    ws_sig.append(["SIG_ENTRY", "LNK_01", 100.0, "NOMINAL", "MAIN"])
    ws_sig.append(["SIG_EXIT", "LNK_02", 500.0, "NOMINAL", "MAIN"])

    ws_blk = wb.create_sheet("Blocks")
    ws_blk.append(["BLOCK_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID"])
    ws_blk.append(["BLK_01", "LNK_01", 0.0, 1000.0, "SIG_ENTRY", "SIG_EXIT"])

    # Both empty; LINK_SEQUENCE with arbitrary spacing
    ws_rte = wb.create_sheet("InterlockingRoutes")
    ws_rte.append(["ROUTE_ID", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID", "LINK_SEQUENCE", "PROTECTED_BLOCKS", "CONFLICTING_ROUTES"])
    ws_rte.append(["RTE_01", "SIG_ENTRY", "SIG_EXIT", "  LNK_01  ,   LNK_02  ", "   ", ""])

    wb.save(sig_file)

    importer = ExcelImporter()
    report = ValidationReport()
    raw_wb = importer.load_workbook(sig_file, report)

    assert raw_wb is not None
    assert not report.has_critical
    routes = raw_wb.get_rows("InterlockingRoutes")
    assert len(routes) == 1
    route = routes[0]
    assert route["link_sequence"] == ["LNK_01", "LNK_02"]
    assert route["protected_blocks"] == []
    assert route["conflicting_routes"] == []


@pytest.mark.integration
def test_simulated_pandas_nan_and_none_in_validator_and_converter():
    """Verify validator and converter tolerate None and NaN in list fields without throwing TypeError."""
    sig_raw = RawWorkbookData(
        "SIGNALLING",
        Path("test_sig.xlsx"),
        {"WORKBOOK_TYPE": "SIGNALLING", "SCHEMA_VERSION": "1.0.0"},
    )
    sig_raw.worksheets["SignallingSystem"] = [
        {"system_type": "FIXED_BLOCK", "aspect_model": "THREE_ASPECT", "overlap_m": 50.0, "release_delay_s": 4.0}
    ]
    sig_raw.worksheets["Signals"] = [
        {"signal_id": "SIG_01", "link_id": "LNK_01", "offset_m": 100.0, "orientation": "FORWARD", "signal_type": "MAIN"},
        {"signal_id": "SIG_02", "link_id": "LNK_02", "offset_m": 200.0, "orientation": "FORWARD", "signal_type": "MAIN"},
    ]
    sig_raw.worksheets["Blocks"] = [
        {"block_id": "BLK_01", "link_id": "LNK_01", "start_offset_m": 0.0, "end_offset_m": 1000.0, "entry_signal_id": "SIG_01", "exit_signal_id": "SIG_02"}
    ]
    # Simulate rows with raw None and NaN
    sig_raw.worksheets["InterlockingRoutes"] = [
        {
            "route_id": "RTE_01",
            "entry_signal_id": "SIG_01",
            "exit_signal_id": "SIG_02",
            "link_sequence": ["LNK_01", "LNK_02"],
            "protected_blocks": None,  # None
            "conflicting_routes": float("nan"),  # NaN
        }
    ]

    infra_raw = RawWorkbookData(
        "INFRASTRUCTURE",
        Path("test_infra.xlsx"),
        {"WORKBOOK_TYPE": "INFRASTRUCTURE", "SCHEMA_VERSION": "1.0.0"},
    )
    infra_raw.worksheets["Nodes"] = [
        {"node_id": "N1", "node_type": "ENDPOINT"},
        {"node_id": "N2", "node_type": "JUNCTION"},
        {"node_id": "N3", "node_type": "ENDPOINT"},
    ]
    infra_raw.worksheets["Tracks"] = [{"track_id": "TRK_01"}]
    infra_raw.worksheets["TrackLinks"] = [
        {"link_id": "LNK_01", "track_id": "TRK_01", "start_node_id": "N1", "end_node_id": "N2", "length_m": 1000.0, "max_speed_ms": 30.0},
        {"link_id": "LNK_02", "track_id": "TRK_01", "start_node_id": "N2", "end_node_id": "N3", "length_m": 1000.0, "max_speed_ms": 30.0},
    ]

    report = ValidationReport()
    validator = DatasetValidator(report)
    # Must NOT raise 'NoneType' object is not iterable
    validator.validate_all({"SIGNALLING": sig_raw, "INFRASTRUCTURE": infra_raw})

    # Must NOT raise during canonical conversion
    canonical = convert_raw_to_canonical({"SIGNALLING": sig_raw, "INFRASTRUCTURE": infra_raw})
    assert len(canonical.signalling.routes) == 1
    route = canonical.signalling.routes[0]
    assert route.route_id == "RTE_01"
    assert route.protected_blocks == []
    assert route.conflicting_route_ids == []


@pytest.mark.integration
def test_user_scenario_447_links_empty_protected_and_conflicting(tmp_path: Path):
    """Verify user's exact scenario:
    One route (RTE_01) with a fully populated LINK_SEQUENCE of 447 links,
    and empty PROTECTED_BLOCKS and CONFLICTING_ROUTES cells.
    """
    n_links = 447
    link_ids = [f"LNK_{i:04d}" for i in range(1, n_links + 1)]
    link_seq_str = ", ".join(link_ids)

    # 1. Build Signalling workbook
    sig_file = tmp_path / "02_Signalling.xlsx"
    wb_sig = openpyxl.Workbook()

    meta_sig = wb_sig.active
    meta_sig.title = "METADATA"
    meta_sig.append(["PROPERTY", "VALUE"])
    meta_sig.append(["WORKBOOK_TYPE", "SIGNALLING"])
    meta_sig.append(["SCHEMA_VERSION", "1.0.0"])

    ws_sys = wb_sig.create_sheet("SignallingSystem")
    ws_sys.append(["TECHNOLOGY_TYPE", "ASPECT_MODEL", "DEFAULT_OVERLAP_M", "SIGHTING_TIME_S", "ROUTE_SETUP_TIME_S", "RELEASE_DELAY_S"])
    ws_sys.append(["GENERIC_FIXED_BLOCK_ENGINEERING_MODEL", "THREE_ASPECT", 50.0, 8.0, 3.0, 4.0])

    ws_sig = wb_sig.create_sheet("Signals")
    ws_sig.append(["SIGNAL_ID", "LINK_ID", "OFFSET_M", "DIRECTION", "SIGNAL_TYPE"])
    ws_sig.append(["SIG_START", link_ids[0], 0.0, "NOMINAL", "MAIN"])
    ws_sig.append(["SIG_END", link_ids[-1], 100.0, "NOMINAL", "MAIN"])

    ws_blk = wb_sig.create_sheet("Blocks")
    ws_blk.append(["BLOCK_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID"])
    ws_blk.append(["BLK_01", link_ids[0], 0.0, 100.0, "SIG_START", "SIG_END"])

    ws_rte = wb_sig.create_sheet("InterlockingRoutes")
    ws_rte.append(["ROUTE_ID", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID", "LINK_SEQUENCE", "PROTECTED_BLOCKS", "CONFLICTING_ROUTES"])
    # Empty cells for PROTECTED_BLOCKS and CONFLICTING_ROUTES
    ws_rte.append(["RTE_01", "SIG_START", "SIG_END", link_seq_str, None, None])

    wb_sig.save(sig_file)

    # 2. Build Infrastructure workbook with 447 links and 448 nodes
    infra_file = tmp_path / "01_Infrastructure.xlsx"
    wb_inf = openpyxl.Workbook()

    meta_inf = wb_inf.active
    meta_inf.title = "METADATA"
    meta_inf.append(["PROPERTY", "VALUE"])
    meta_inf.append(["WORKBOOK_TYPE", "INFRASTRUCTURE"])
    meta_inf.append(["SCHEMA_VERSION", "1.0.0"])

    ws_nodes = wb_inf.create_sheet("Nodes")
    ws_nodes.append(["NODE_ID", "NODE_TYPE"])
    for i in range(1, n_links + 2):
        ws_nodes.append([f"ND_{i:04d}", "ENDPOINT" if i in (1, n_links + 1) else "SWITCH"])

    ws_trks = wb_inf.create_sheet("Tracks")
    ws_trks.append(["TRACK_ID", "DESCRIPTION", "DIRECTIONALITY"])
    ws_trks.append(["TRK_01", "Main Line", "NOMINAL"])

    ws_lnks = wb_inf.create_sheet("TrackLinks")
    ws_lnks.append(["LINK_ID", "TRACK_ID", "START_NODE_ID", "END_NODE_ID", "LENGTH_M", "MAX_SPEED_KMH", "GRADIENT_PERMIL"])
    for i in range(1, n_links + 1):
        ws_lnks.append([f"LNK_{i:04d}", "TRK_01", f"ND_{i:04d}", f"ND_{i+1:04d}", 100.0, 120.0, 0.0])

    wb_inf.save(infra_file)

    # Ingest using ExcelImporter
    importer = ExcelImporter()
    report = ValidationReport()

    raw_sig = importer.load_workbook(sig_file, report)
    raw_inf = importer.load_workbook(infra_file, report)

    assert raw_sig is not None, "Signalling workbook failed to load"
    assert raw_inf is not None, "Infrastructure workbook failed to load"
    assert not report.has_critical, f"Critical findings during import: {[f.message for f in report.findings]}"

    # Validate dataset
    validator = DatasetValidator(report)
    validator.validate_all({"SIGNALLING": raw_sig, "INFRASTRUCTURE": raw_inf})
    # Filter out Level 5 incomplete workbooks check since we are only validating SIGNALLING + INFRASTRUCTURE
    relevant_errors = [
        f for f in report.findings
        if f.error_code != "ERR_SIMULATION_READINESS_INCOMPLETE" and f.severity.value in ("ERROR", "CRITICAL")
    ]
    assert len(relevant_errors) == 0, f"Validation errors: {[f.message for f in relevant_errors]}"

    # Convert to canonical model
    canonical = convert_raw_to_canonical({"SIGNALLING": raw_sig, "INFRASTRUCTURE": raw_inf})
    assert len(canonical.signalling.routes) == 1
    route = canonical.signalling.routes[0]
    assert route.route_id == "RTE_01"
    assert len(route.link_sequence) == 447
    assert route.link_sequence == link_ids
    assert route.protected_blocks == []
    assert route.conflicting_route_ids == []


@pytest.mark.unit
def test_fully_populated_route_preserves_behavior():
    """Verify existing behavior is strictly preserved for fully populated routes."""
    col_def = ColumnDef(
        name="PROTECTED_BLOCKS",
        canonical_name="protected_blocks",
        field_type=FieldType.STRING_LIST,
        required=False,
    )
    report = ValidationReport()
    loc_ctx = {"workbook": "test.xlsx", "worksheet": "InterlockingRoutes", "row": 2}

    val, ok = _parse_cell_value("BLK_01, BLK_02, BLK_03", col_def, report, loc_ctx)
    assert ok is True
    assert val == ["BLK_01", "BLK_02", "BLK_03"]
