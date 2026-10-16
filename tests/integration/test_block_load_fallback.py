"""Regression tests for block loading fallback removal.

RHS-CORRECTION-BLOCK-LOAD-FALLBACK-001:
  - Task 2 (Option A): Fail loudly with ERR_SIG_BLOCKS_MISSING when zero blocks loaded.
  - Task 4 & 5: Ensure runtime resource list uses user-defined blocks and never synthesizes BLK_LNK_*.
"""

from pathlib import Path
import openpyxl
import pytest

from headway.core.exceptions import SignallingBlocksMissingError
from headway.data.converter import convert_raw_to_canonical
from headway.data.importer import ExcelImporter
from headway.data.validation import ValidationReport
from headway.infrastructure.direction import RunningDirection
from headway.ui.simulation_runner import SimulationRunner


def _create_minimal_infrastructure_wb(path: Path, link_ids: list[str]) -> Path:
    wb = openpyxl.Workbook()
    meta = wb.active
    meta.title = "METADATA"
    meta.append(["PROPERTY", "VALUE"])
    meta.append(["WORKBOOK_TYPE", "INFRASTRUCTURE"])
    meta.append(["SCHEMA_VERSION", "1.0.0"])

    ws_nodes = wb.create_sheet("Nodes")
    ws_nodes.append(["NODE_ID", "NODE_TYPE"])
    for i in range(1, len(link_ids) + 2):
        ws_nodes.append([f"ND_{i:04d}", "ENDPOINT" if i in (1, len(link_ids) + 1) else "SWITCH"])

    ws_trks = wb.create_sheet("Tracks")
    ws_trks.append(["TRACK_ID", "DESCRIPTION", "DIRECTIONALITY"])
    ws_trks.append(["TRK_01", "Main Line", "NOMINAL"])

    ws_lnks = wb.create_sheet("TrackLinks")
    ws_lnks.append(["LINK_ID", "TRACK_ID", "START_NODE_ID", "END_NODE_ID", "LENGTH_M", "MAX_SPEED_KMH", "GRADIENT_PERMIL"])
    for i, lnk in enumerate(link_ids, start=1):
        ws_lnks.append([lnk, "TRK_01", f"ND_{i:04d}", f"ND_{i+1:04d}", 1000.0, 120.0, 0.0])

    for s_name in ["Stations", "Platforms", "StoppingPoints", "Tunnels", "TVSSections", "TVSGeometry", "SharedResourceGroups"]:
        ws = wb.create_sheet(s_name)
        if s_name == "Stations":
            ws.append(["STATION_ID", "NAME"])
            ws.append(["STN_01", "Station 1"])
        elif s_name == "Platforms":
            ws.append(["PLATFORM_ID", "STATION_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "LENGTH_M"])
            ws.append(["PLT_01", "STN_01", link_ids[0], 0.0, 200.0, 200.0])
        elif s_name == "StoppingPoints":
            ws.append(["STOPPING_POINT_ID", "PLATFORM_ID", "LINK_ID", "OFFSET_M"])
            ws.append(["SP_01", "PLT_01", link_ids[0], 100.0])
        elif s_name == "Tunnels":
            ws.append(["TUNNEL_ID", "NAME"])
        elif s_name == "TVSSections":
            ws.append(["TVS_ID", "TUNNEL_ID", "TRACK_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "RELEASE_DELAY_S"])
        elif s_name == "TVSGeometry":
            ws.append(["TVS_ID", "SEQUENCE", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M"])
        elif s_name == "SharedResourceGroups":
            ws.append(["GROUP_ID", "DESCRIPTION", "RESOURCE_IDS"])

    wb.save(path)
    return path


def _create_minimal_rolling_stock_wb(path: Path) -> Path:
    wb = openpyxl.Workbook()
    meta = wb.active
    meta.title = "METADATA"
    meta.append(["PROPERTY", "VALUE"])
    meta.append(["WORKBOOK_TYPE", "ROLLING_STOCK"])
    meta.append(["SCHEMA_VERSION", "1.0.0"])

    ws_tt = wb.create_sheet("TrainTypes")
    ws_tt.append([
        "TRAIN_TYPE_ID", "DESCRIPTION", "LENGTH_M", "MASS_EMPTY_T", "MASS_LOADED_T",
        "ROTATING_MASS_FACTOR", "MAX_SPEED_KMH", "MAX_ACCELERATION_MS2", "SERVICE_DECELERATION_MS2",
        "EMERGENCY_DECELERATION_MS2", "TRACTION_MODEL_TYPE", "INSTALLED_POWER_KW",
        "MAX_TRACTIVE_EFFORT_KN", "BRAKING_MODEL_TYPE", "DAVIS_A_N", "DAVIS_B_NS_PER_M",
        "DAVIS_C_NS2_PER_M2",
    ])
    ws_tt.append([
        "TRN_01", "Standard Passenger Train", 120.0, 180.0, 220.0,
        0.08, 120.0, 1.0, 0.9,
        1.2, "SIMPLIFIED_POWER_FORCE", 3000.0,
        200.0, "CONSTANT_DECELERATION", 1500.0, 25.0,
        0.35,
    ])
    wb.save(path)
    return path


def test_three_user_blocks_over_three_links_runtime_resources(tmp_path: Path):
    """Test 5a: Loads 3 user-defined blocks over 3 links, verifies exactly 3 user block IDs, no BLK_LNK_*."""
    link_ids = ["LNK_01", "LNK_02", "LNK_03"]
    infra_file = _create_minimal_infrastructure_wb(tmp_path / "01_Infrastructure.xlsx", link_ids)
    rs_file = _create_minimal_rolling_stock_wb(tmp_path / "03_RollingStock.xlsx")

    # Signalling workbook with 3 user blocks
    sig_file = tmp_path / "02_Signalling.xlsx"
    wb = openpyxl.Workbook()
    meta = wb.active
    meta.title = "METADATA"
    meta.append(["PROPERTY", "VALUE"])
    meta.append(["WORKBOOK_TYPE", "SIGNALLING"])
    meta.append(["SCHEMA_VERSION", "1.0.0"])

    ws_sys = wb.create_sheet("SignallingSystem")
    ws_sys.append(["TECHNOLOGY_TYPE", "ASPECT_MODEL", "DEFAULT_OVERLAP_M", "SIGHTING_TIME_S", "ROUTE_SETUP_TIME_S", "RELEASE_DELAY_S"])
    ws_sys.append(["GENERIC_FIXED_BLOCK_ENGINEERING_MODEL", "THREE_ASPECT", 50.0, 8.0, 3.0, 4.0])

    ws_sig = wb.create_sheet("Signals")
    ws_sig.append(["SIGNAL_ID", "LINK_ID", "OFFSET_M", "DIRECTION", "SIGNAL_TYPE"])
    ws_sig.append(["SIG_01", "LNK_01", 100.0, "NOMINAL", "MAIN"])
    ws_sig.append(["SIG_02", "LNK_02", 100.0, "NOMINAL", "MAIN"])
    ws_sig.append(["SIG_03", "LNK_03", 100.0, "NOMINAL", "MAIN"])

    ws_blk = wb.create_sheet("Blocks")
    ws_blk.append(["BLOCK_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID"])
    user_block_ids = ["BLK_USER_01", "BLK_USER_02", "BLK_USER_03"]
    for idx, (b_id, l_id) in enumerate(zip(user_block_ids, link_ids), start=1):
        ws_blk.append([b_id, l_id, 0.0, 1000.0, f"SIG_{idx:02d}", None])

    ws_rte = wb.create_sheet("InterlockingRoutes")
    ws_rte.append(["ROUTE_ID", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID", "LINK_SEQUENCE", "PROTECTED_BLOCKS", "CONFLICTING_ROUTES"])
    ws_rte.append(["RTE_01", "SIG_01", "SIG_03", "LNK_01, LNK_02, LNK_03", "BLK_USER_01, BLK_USER_02, BLK_USER_03", ""])
    wb.save(sig_file)

    # Ingest and convert
    importer = ExcelImporter()
    report = ValidationReport()
    raw_sig = importer.load_workbook(sig_file, report)
    raw_inf = importer.load_workbook(infra_file, report)
    raw_rs = importer.load_workbook(rs_file, report)

    project = convert_raw_to_canonical({"SIGNALLING": raw_sig, "INFRASTRUCTURE": raw_inf, "ROLLING_STOCK": raw_rs})

    # Run analysis pipeline
    pkg, doc = SimulationRunner.run_simulation_pipeline(project=project, running_direction=RunningDirection.FORWARD)

    # Assert runtime resource list contains exactly the 3 user-defined block IDs
    resource_ids = [b.resource_id for b in pkg.blocking_intervals]
    assert resource_ids == user_block_ids
    assert len(resource_ids) == 3

    # Assert that no resource has an ID matching the pattern BLK_LNK_*
    assert not any("BLK_LNK_" in rid for rid in resource_ids)
    assert not any(rid.startswith("BLK_LNK") for rid in resource_ids)
    assert not any("BLK_LNK_" in rec.resource_id for rec in pkg.metadata.get("provenance_records", []))


def test_zero_blocks_raises_err_sig_blocks_missing(tmp_path: Path):
    """Test 5b: Loading a workbook with zero blocks raises ERR_SIG_BLOCKS_MISSING."""
    link_ids = ["LNK_01", "LNK_02"]
    infra_file = _create_minimal_infrastructure_wb(tmp_path / "01_Infrastructure.xlsx", link_ids)
    rs_file = _create_minimal_rolling_stock_wb(tmp_path / "03_RollingStock.xlsx")

    # Signalling workbook with ZERO blocks
    sig_file = tmp_path / "02_Signalling.xlsx"
    wb = openpyxl.Workbook()
    meta = wb.active
    meta.title = "METADATA"
    meta.append(["PROPERTY", "VALUE"])
    meta.append(["WORKBOOK_TYPE", "SIGNALLING"])
    meta.append(["SCHEMA_VERSION", "1.0.0"])

    ws_sys = wb.create_sheet("SignallingSystem")
    ws_sys.append(["TECHNOLOGY_TYPE", "ASPECT_MODEL", "DEFAULT_OVERLAP_M", "SIGHTING_TIME_S", "ROUTE_SETUP_TIME_S", "RELEASE_DELAY_S"])
    ws_sys.append(["GENERIC_FIXED_BLOCK_ENGINEERING_MODEL", "THREE_ASPECT", 50.0, 8.0, 3.0, 4.0])

    ws_sig = wb.create_sheet("Signals")
    ws_sig.append(["SIGNAL_ID", "LINK_ID", "OFFSET_M", "DIRECTION", "SIGNAL_TYPE"])
    ws_sig.append(["SIG_01", "LNK_01", 100.0, "NOMINAL", "MAIN"])

    ws_blk = wb.create_sheet("Blocks")
    ws_blk.append(["BLOCK_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID"])
    # Zero block rows added

    ws_rte = wb.create_sheet("InterlockingRoutes")
    ws_rte.append(["ROUTE_ID", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID", "LINK_SEQUENCE", "PROTECTED_BLOCKS", "CONFLICTING_ROUTES"])
    ws_rte.append(["RTE_01", "SIG_01", "SIG_01", "LNK_01", "", ""])
    wb.save(sig_file)

    importer = ExcelImporter()
    report = ValidationReport()
    raw_sig = importer.load_workbook(sig_file, report)
    raw_inf = importer.load_workbook(infra_file, report)
    raw_rs = importer.load_workbook(rs_file, report)

    project = convert_raw_to_canonical({"SIGNALLING": raw_sig, "INFRASTRUCTURE": raw_inf, "ROLLING_STOCK": raw_rs})

    with pytest.raises(SignallingBlocksMissingError) as exc_info:
        SimulationRunner.run_simulation_pipeline(project=project, running_direction=RunningDirection.FORWARD)

    assert "ERR_SIG_BLOCKS_MISSING" in str(exc_info.value)
    assert exc_info.value.error_code == "ERR_SIG_BLOCKS_MISSING"
    assert "No signalling blocks were loaded from the Signalling workbook" in str(exc_info.value)


def test_user_data_shape_447_links_92_blocks(tmp_path: Path):
    """Test 4: User scenario with 448 nodes, 447 links, 91 signals, and 92 blocks BLK_UP_001..092."""
    n_links = 447
    link_ids = [f"LNK_{i:04d}" for i in range(1, n_links + 1)]
    infra_file = _create_minimal_infrastructure_wb(tmp_path / "01_Infrastructure.xlsx", link_ids)
    rs_file = _create_minimal_rolling_stock_wb(tmp_path / "03_RollingStock.xlsx")

    # Signalling workbook with 92 blocks and 91 signals
    sig_file = tmp_path / "02_Signalling.xlsx"
    wb = openpyxl.Workbook()
    meta = wb.active
    meta.title = "METADATA"
    meta.append(["PROPERTY", "VALUE"])
    meta.append(["WORKBOOK_TYPE", "SIGNALLING"])
    meta.append(["SCHEMA_VERSION", "1.0.0"])

    ws_sys = wb.create_sheet("SignallingSystem")
    ws_sys.append(["TECHNOLOGY_TYPE", "ASPECT_MODEL", "DEFAULT_OVERLAP_M", "SIGHTING_TIME_S", "ROUTE_SETUP_TIME_S", "RELEASE_DELAY_S"])
    ws_sys.append(["GENERIC_FIXED_BLOCK_ENGINEERING_MODEL", "THREE_ASPECT", 50.0, 8.0, 3.0, 4.0])

    ws_sig = wb.create_sheet("Signals")
    ws_sig.append(["SIGNAL_ID", "LINK_ID", "OFFSET_M", "DIRECTION", "SIGNAL_TYPE"])
    for i in range(1, 92):
        sig_id = f"SIG_{i:03d}"
        lnk_idx = min(i * 4, n_links - 1)
        ws_sig.append([sig_id, link_ids[lnk_idx], 50.0, "NOMINAL", "MAIN"])

    ws_blk = wb.create_sheet("Blocks")
    ws_blk.append(["BLOCK_ID", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID"])
    ws_geom = wb.create_sheet("BlockGeometry")
    ws_geom.append(["BLOCK_ID", "SEQUENCE", "LINK_ID", "START_OFFSET_M", "END_OFFSET_M"])

    for i in range(1, 93):
        blk_id = f"BLK_UP_{i:03d}"
        entry_sig = f"SIG_{i:03d}" if i <= 91 else None
        exit_sig = f"SIG_{i+1:03d}" if i < 91 else None
        start_lnk_idx = (i - 1) * 4
        end_lnk_idx = min(i * 4, n_links)
        ws_blk.append([blk_id, link_ids[start_lnk_idx], 0.0, 100.0, entry_sig, exit_sig])
        seq = 1
        for l_idx in range(start_lnk_idx, end_lnk_idx):
            ws_geom.append([blk_id, seq, link_ids[l_idx], 0.0, 100.0])
            seq += 1

    ws_rte = wb.create_sheet("InterlockingRoutes")
    ws_rte.append(["ROUTE_ID", "ENTRY_SIGNAL_ID", "EXIT_SIGNAL_ID", "LINK_SEQUENCE", "PROTECTED_BLOCKS", "CONFLICTING_ROUTES"])
    ws_rte.append(["RTE_01", "SIG_001", "SIG_091", ", ".join(link_ids), "", ""])
    wb.save(sig_file)

    importer = ExcelImporter()
    report = ValidationReport()
    raw_sig = importer.load_workbook(sig_file, report)
    raw_inf = importer.load_workbook(infra_file, report)
    raw_rs = importer.load_workbook(rs_file, report)

    project = convert_raw_to_canonical({"SIGNALLING": raw_sig, "INFRASTRUCTURE": raw_inf, "ROLLING_STOCK": raw_rs})

    pkg, doc = SimulationRunner.run_simulation_pipeline(project=project, running_direction=RunningDirection.FORWARD)

    # Confirm 92 blocks loaded
    assert len(project.signalling.blocks) == 92
    resource_ids = [b.resource_id for b in pkg.blocking_intervals]
    assert len(resource_ids) == 92

    # First 5 block IDs
    assert resource_ids[:5] == ["BLK_UP_001", "BLK_UP_002", "BLK_UP_003", "BLK_UP_004", "BLK_UP_005"]
    # Last 5 block IDs
    assert resource_ids[-5:] == ["BLK_UP_088", "BLK_UP_089", "BLK_UP_090", "BLK_UP_091", "BLK_UP_092"]

    # No synthetic IDs
    assert not any("BLK_LNK_" in rid for rid in resource_ids)
    assert not any("BLK_LNK_" in rec.resource_id for rec in pkg.metadata.get("provenance_records", []))
