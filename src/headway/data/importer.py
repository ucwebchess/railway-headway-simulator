"""Excel workbook importer and raw worksheet extractor.

Strictly satisfies P01-IMP-001 through P01-IMP-009:
- Reads .xlsx workbooks using openpyxl (rejects macro .xlsm and invalid formats).
- Detects workbook type via METADATA sheet, not filenames.
- Reads cached formula values without executing macros (data_only=True).
- Uses the authoritative schema registry to parse columns and types.
- Distinguishes example rows (e.g. 'EXAMPLE_*') from project data.
- Emits structured ValidationFindings with exact workbook, worksheet, row, and column context.
"""

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union
import openpyxl

from headway.core.exceptions import DataValidationError
from headway.data.excel_registry import (
    ALL_WORKBOOK_DEFS,
    ColumnDef,
    FieldType,
    WorkbookDef,
    WorksheetDef,
    get_workbook_def,
)
from headway.data.validation import Severity, ValidationReport


class RawWorkbookData:
    """Container for parsed raw worksheet records from a single Excel workbook."""

    def __init__(self, workbook_type: str, file_path: Path, metadata: Dict[str, str]) -> None:
        self.workbook_type: str = workbook_type
        self.file_path: Path = file_path
        self.metadata: Dict[str, str] = metadata
        self.worksheets: Dict[str, List[Dict[str, Any]]] = {}

    def get_rows(self, worksheet_name: str) -> List[Dict[str, Any]]:
        return self.worksheets.get(worksheet_name, [])


def _parse_metadata(sheet: openpyxl.worksheet.worksheet.Worksheet) -> Dict[str, str]:
    """Extract key-value pairs from METADATA worksheet."""
    meta: Dict[str, str] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip().upper()
        val = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        meta[key] = val
    return meta


def _is_example_or_subheader_row(row_idx: int, row_values: List[Any], columns: List[ColumnDef]) -> bool:
    """Determine whether a row is a unit subheader or instructional example row."""
    if not any(row_values):
        return True  # Empty row

    # Subheader row typically has brackets like [m], [km/h], or (Required)/(Optional)
    first_val = str(row_values[0]).strip() if row_values[0] is not None else ""
    if row_idx == 2 and (first_val.startswith("[") or first_val.startswith("(") or "Required" in first_val):
        return True

    # Check for EXAMPLE_* identifier prefix in the first column
    if first_val.upper().startswith("EXAMPLE_") or first_val.upper().startswith("SAMPLE_"):
        return True

    return False


def _parse_cell_value(
    raw_val: Any,
    col_def: ColumnDef,
    report: ValidationReport,
    loc_ctx: Dict[str, Any],
) -> Tuple[Any, bool]:
    """Parse and validate a single cell value according to ColumnDef."""
    if raw_val is None or str(raw_val).strip() == "":
        if col_def.required:
            report.add_finding(
                error_code="ERR_FIELD_REQUIRED_MISSING",
                severity=Severity.ERROR,
                message=f"Mandatory column '{col_def.name}' is blank.",
                column=col_def.name,
                recommendation=f"Provide a valid {col_def.field_type.value} value.",
                **loc_ctx,
            )
            return None, False
        return col_def.default_value, True

    # Clean string representation
    s_val = str(raw_val).strip()

    try:
        if col_def.field_type == FieldType.STRING:
            # Check allowed enum values if defined
            if col_def.allowed_values and s_val not in col_def.allowed_values:
                # Case-insensitive match check
                matched = False
                for allowed in col_def.allowed_values:
                    if s_val.upper() == allowed.upper():
                        s_val = allowed
                        matched = True
                        break
                if not matched:
                    report.add_finding(
                        error_code="ERR_FIELD_ENUM_INVALID",
                        severity=Severity.ERROR,
                        message=f"Invalid value '{s_val}' for column '{col_def.name}'. Allowed values: {col_def.allowed_values}",
                        column=col_def.name,
                        recommendation=f"Select one of: {', '.join(col_def.allowed_values)}",
                        **loc_ctx,
                    )
                    return None, False
            return s_val, True

        elif col_def.field_type == FieldType.FLOAT:
            f_val = float(raw_val)
            return f_val, True

        elif col_def.field_type == FieldType.INTEGER:
            i_val = int(round(float(raw_val)))
            return i_val, True

        elif col_def.field_type == FieldType.BOOLEAN:
            if isinstance(raw_val, bool):
                return raw_val, True
            if s_val.lower() in {"true", "1", "yes", "y", "t"}:
                return True, True
            if s_val.lower() in {"false", "0", "no", "n", "f"}:
                return False, True
            report.add_finding(
                error_code="ERR_FIELD_BOOLEAN_INVALID",
                severity=Severity.ERROR,
                message=f"Cannot parse '{s_val}' as boolean in column '{col_def.name}'.",
                column=col_def.name,
                recommendation="Use TRUE or FALSE.",
                **loc_ctx,
            )
            return None, False

        elif col_def.field_type == FieldType.STRING_LIST:
            # Comma or semicolon separated list
            parts = [p.strip() for p in s_val.replace(";", ",").split(",") if p.strip()]
            return parts, True

    except (ValueError, TypeError) as err:
        report.add_finding(
            error_code="ERR_FIELD_TYPE_INVALID",
            severity=Severity.ERROR,
            message=f"Value '{raw_val}' in column '{col_def.name}' cannot be parsed as {col_def.field_type.value}: {err}",
            column=col_def.name,
            recommendation=f"Enter a valid numeric/text value conforming to {col_def.field_type.value}.",
            **loc_ctx,
        )
        return None, False

    return s_val, True


class ExcelImporter:
    """Reads and validates Excel workbooks against the authoritative registry."""

    def __init__(self) -> None:
        pass

    def load_workbook(self, file_path: Union[str, Path], report: ValidationReport) -> Optional[RawWorkbookData]:
        """Load and parse an Excel workbook file into RawWorkbookData.

        Performs Level 1 (Structural) and Level 2 (Field) validation.
        """
        path = Path(file_path)
        wb_name = path.name

        # Level 1 check: file extension
        if not path.exists():
            report.add_finding(
                error_code="ERR_STRUCTURAL_FILE_NOT_FOUND",
                severity=Severity.CRITICAL,
                message=f"Excel file '{path}' does not exist.",
                workbook=wb_name,
            )
            return None

        if path.suffix.lower() == ".xlsm":
            report.add_finding(
                error_code="ERR_SECURITY_MACRO_REJECTED",
                severity=Severity.CRITICAL,
                message=f"Macro-enabled workbook '{wb_name}' is strictly prohibited for security reasons.",
                workbook=wb_name,
                recommendation="Convert file to standard .xlsx format without macros.",
            )
            return None

        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            report.add_finding(
                error_code="ERR_STRUCTURAL_FILE_TYPE",
                severity=Severity.CRITICAL,
                message=f"Unsupported file format '{path.suffix}'. Only .xlsx is supported.",
                workbook=wb_name,
                recommendation="Provide a valid Microsoft Excel .xlsx workbook.",
            )
            return None

        # Load with openpyxl
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except Exception as err:
            report.add_finding(
                error_code="ERR_STRUCTURAL_CORRUPT_FILE",
                severity=Severity.CRITICAL,
                message=f"Failed to open Excel workbook '{wb_name}': {err}",
                workbook=wb_name,
                recommendation="Ensure the file is a valid, uncorrupted Excel workbook.",
            )
            return None

        # Level 1 check: METADATA worksheet
        if "METADATA" not in wb.sheetnames:
            report.add_finding(
                error_code="ERR_STRUCTURAL_METADATA_MISSING",
                severity=Severity.CRITICAL,
                message=f"Workbook '{wb_name}' lacks required 'METADATA' worksheet.",
                workbook=wb_name,
                worksheet="METADATA",
                recommendation="Add METADATA worksheet declaring WORKBOOK_TYPE and SCHEMA_VERSION.",
            )
            return None

        metadata = _parse_metadata(wb["METADATA"])
        wb_type = metadata.get("WORKBOOK_TYPE", "").strip().upper()
        schema_version = metadata.get("SCHEMA_VERSION", "").strip()

        if not wb_type:
            report.add_finding(
                error_code="ERR_STRUCTURAL_WORKBOOK_TYPE_MISSING",
                severity=Severity.CRITICAL,
                message=f"Workbook '{wb_name}' METADATA sheet missing 'WORKBOOK_TYPE'.",
                workbook=wb_name,
                worksheet="METADATA",
            )
            return None

        wb_def = get_workbook_def(wb_type)
        if wb_def is None:
            report.add_finding(
                error_code="ERR_STRUCTURAL_UNKNOWN_TYPE",
                severity=Severity.CRITICAL,
                message=f"Unknown WORKBOOK_TYPE '{wb_type}' in '{wb_name}'. Recognized: {list(ALL_WORKBOOK_DEFS.keys())}",
                workbook=wb_name,
                worksheet="METADATA",
            )
            return None

        if schema_version != "1.0.0":
            report.add_finding(
                error_code="ERR_STRUCTURAL_UNSUPPORTED_VERSION",
                severity=Severity.ERROR,
                message=f"Unsupported SCHEMA_VERSION '{schema_version}' in '{wb_name}'. Expected: '1.0.0'.",
                workbook=wb_name,
                worksheet="METADATA",
                recommendation="Upgrade workbook to Schema Version 1.0.0.",
            )

        raw_wb = RawWorkbookData(workbook_type=wb_type, file_path=path, metadata=metadata)

        # Level 1 check: Required worksheets
        for ws_def in wb_def.worksheets:
            if ws_def.name == "METADATA":
                continue

            if ws_def.name not in wb.sheetnames:
                if ws_def.required:
                    report.add_finding(
                        error_code="ERR_STRUCTURAL_WORKSHEET_MISSING",
                        severity=Severity.CRITICAL,
                        message=f"Required worksheet '{ws_def.name}' is missing in '{wb_name}'.",
                        workbook=wb_name,
                        worksheet=ws_def.name,
                        recommendation=f"Add worksheet '{ws_def.name}' to workbook.",
                    )
                continue

            sheet = wb[ws_def.name]

            # Level 1 check: Header columns (Row 1)
            row_1_cells = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1), [])]
            raw_headers = [str(c).strip() for c in row_1_cells if c is not None and str(c).strip()]

            # Duplicate column detection
            seen_cols = set()
            for col_h in raw_headers:
                if col_h.upper() in seen_cols:
                    report.add_finding(
                        error_code="ERR_STRUCTURAL_DUPLICATE_COLUMN",
                        severity=Severity.ERROR,
                        message=f"Duplicate column header '{col_h}' found in worksheet '{ws_def.name}'.",
                        workbook=wb_name,
                        worksheet=ws_def.name,
                        column=col_h,
                    )
                seen_cols.add(col_h.upper())

            # Map column index to ColumnDef
            col_map: Dict[int, ColumnDef] = {}
            for col_idx, raw_col_name in enumerate(row_1_cells, start=1):
                if raw_col_name is None:
                    continue
                clean_name = str(raw_col_name).strip()
                # Find matching ColumnDef (case-insensitive)
                for def_col in ws_def.columns:
                    if def_col.name.upper() == clean_name.upper():
                        col_map[col_idx] = def_col
                        break

            # Verify required columns exist
            for def_col in ws_def.columns:
                if def_col.required and not any(c.name.upper() == def_col.name.upper() for c in col_map.values()):
                    report.add_finding(
                        error_code="ERR_STRUCTURAL_COLUMN_MISSING",
                        severity=Severity.CRITICAL,
                        message=f"Mandatory column '{def_col.name}' is missing in worksheet '{ws_def.name}'.",
                        workbook=wb_name,
                        worksheet=ws_def.name,
                        column=def_col.name,
                        recommendation=f"Add column '{def_col.name}' with appropriate values.",
                    )

            # Parse data rows (Level 2: Field validation)
            parsed_rows = []
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not row or not any(row):
                    continue

                if _is_example_or_subheader_row(row_idx, list(row), ws_def.columns):
                    continue

                row_dict: Dict[str, Any] = {}
                row_valid = True

                for col_idx, val in enumerate(row, start=1):
                    col_def = col_map.get(col_idx)
                    if not col_def:
                        continue  # Extra unrecognized column

                    first_id = str(row[0]).strip() if len(row) > 0 and row[0] is not None else None
                    loc_ctx = {
                        "workbook": wb_name,
                        "worksheet": ws_def.name,
                        "row": row_idx,
                        "object_id": first_id,
                    }

                    parsed_val, is_valid = _parse_cell_value(val, col_def, report, loc_ctx)
                    if not is_valid and col_def.required:
                        row_valid = False
                    row_dict[col_def.canonical_name] = parsed_val

                # Fill defaults for omitted columns
                for def_col in ws_def.columns:
                    if def_col.canonical_name not in row_dict:
                        if def_col.required:
                            row_valid = False
                        else:
                            row_dict[def_col.canonical_name] = def_col.default_value

                if row_valid and any(v is not None for v in row_dict.values()):
                    parsed_rows.append(row_dict)

            # Check min_rows requirement
            if ws_def.required and ws_def.min_rows > 0 and len(parsed_rows) < ws_def.min_rows:
                report.add_finding(
                    error_code="ERR_STRUCTURAL_INSUFFICIENT_ROWS",
                    severity=Severity.ERROR,
                    message=f"Worksheet '{ws_def.name}' contains {len(parsed_rows)} valid data rows, but at least {ws_def.min_rows} required.",
                    workbook=wb_name,
                    worksheet=ws_def.name,
                    recommendation="Ensure data rows are provided and not mistaken for example rows.",
                )

            raw_wb.worksheets[ws_def.name] = parsed_rows

        return raw_wb
