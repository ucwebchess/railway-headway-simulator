"""Structured validation framework and finding reports for Railway Headway Simulator.

Strictly satisfies RHS-P01-001 § 15 & § 16:
- Standardized validation findings with error code, severity, worksheet, row, column, message, recommendation.
- Severity levels: CRITICAL, ERROR, WARNING, INFO.
- Prevents simulation readiness when CRITICAL or ERROR findings exist.
- Supports exporting validation findings to Excel workbook.
"""

from dataclasses import asdict, dataclass
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill


class Severity(str, Enum):
    CRITICAL = "CRITICAL"
    ERROR = "ERROR"
    WARNING = "WARNING"
    INFO = "INFO"


@dataclass
class ValidationFinding:
    """Individual engineering or data contract validation finding."""

    error_code: str
    severity: Severity
    message: str
    dataset_id: Optional[str] = None
    workbook: Optional[str] = None
    worksheet: Optional[str] = None
    row: Optional[int] = None
    column: Optional[str] = None
    object_id: Optional[str] = None
    recommendation: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        """Convert finding to standard dictionary."""
        d = asdict(self)
        d["severity"] = self.severity.value
        return d


class ValidationReport:
    """Aggregate validation findings container for a dataset or project."""

    def __init__(self, project_id: Optional[str] = None) -> None:
        self.project_id: Optional[str] = project_id
        self.findings: List[ValidationFinding] = []

    def add(self, finding: ValidationFinding) -> None:
        """Append a validation finding."""
        self.findings.append(finding)

    def add_finding(
        self,
        error_code: str,
        severity: Severity,
        message: str,
        dataset_id: Optional[str] = None,
        workbook: Optional[str] = None,
        worksheet: Optional[str] = None,
        row: Optional[int] = None,
        column: Optional[str] = None,
        object_id: Optional[str] = None,
        recommendation: Optional[str] = None,
    ) -> None:
        """Construct and append a validation finding."""
        self.findings.append(
            ValidationFinding(
                error_code=error_code,
                severity=severity,
                message=message,
                dataset_id=dataset_id,
                workbook=workbook,
                worksheet=worksheet,
                row=row,
                column=column,
                object_id=object_id,
                recommendation=recommendation,
            )
        )

    @property
    def has_critical(self) -> bool:
        return any(f.severity == Severity.CRITICAL for f in self.findings)

    @property
    def has_errors(self) -> bool:
        return any(f.severity in {Severity.CRITICAL, Severity.ERROR} for f in self.findings)

    @property
    def is_simulation_ready(self) -> bool:
        """Determines whether inputs satisfy criteria to advance to microscopic simulation."""
        return not self.has_errors

    def summary(self) -> Dict[str, int]:
        """Return counts by severity level."""
        counts = {s.value: 0 for s in Severity}
        for f in self.findings:
            counts[f.severity.value] += 1
        return counts

    def to_list(self) -> List[Dict[str, Any]]:
        """Return list of dictionary findings."""
        return [f.to_dict() for f in self.findings]

    def export_to_excel(self, output_path: Path) -> Path:
        """Export all validation findings to a standardized Excel workbook."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Validation_Findings"
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "ERROR_CODE", "SEVERITY", "WORKBOOK", "WORKSHEET",
            "ROW", "COLUMN", "OBJECT_ID", "MESSAGE", "RECOMMENDATION"
        ]

        # Header styling
        header_fill = PatternFill(start_color="0F2B48", end_color="0F2B48", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font

        # Colors for severity
        sev_colors = {
            Severity.CRITICAL: "FEE2E2",  # Light red
            Severity.ERROR: "FFEDD5",     # Light orange
            Severity.WARNING: "FEF9C3",   # Light yellow
            Severity.INFO: "E0F2FE",      # Light blue
        }

        for r_idx, finding in enumerate(self.findings, start=2):
            row_vals = [
                finding.error_code,
                finding.severity.value,
                finding.workbook or "",
                finding.worksheet or "",
                finding.row if finding.row is not None else "",
                finding.column or "",
                finding.object_id or "",
                finding.message,
                finding.recommendation or "",
            ]
            fill_color = sev_colors.get(finding.severity)
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid") if fill_color else None

            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                if row_fill and c_idx == 2:  # Highlight severity column
                    cell.fill = row_fill
                cell.font = Font(name="Calibri", size=10)

        # Adjust column widths
        for col_letter, width in zip(["A", "B", "C", "D", "E", "F", "G", "H", "I"], [18, 12, 25, 18, 8, 18, 16, 50, 40]):
            ws.column_dimensions[col_letter].width = width

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return output_path
